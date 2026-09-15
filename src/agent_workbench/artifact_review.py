"""Bounded snapshots, review revisions, and process-local consumption bindings.

Extraction contract: pass registry=registry (or validator=...) to preparation.
Executable reviews also require engine-owned step_id and execution_mode. At
dispatch, supply the current run_id on AdapterInput or as a verifier keyword;
never reconstruct the current execution context from an old binding.
The validator implements validate(snapshot, policy); missing validation blocks.
Content rules belong to plugins. CSV/JSON readers and optional XLSX only parse.
Evidence contains opaque IDs and explicitly selected, redacted columns, never
source paths or raw bytes. Directory manifests cover all bounded membership,
but incidental contents are neither opened nor copied.

Stable descriptor reads, before/after membership checks, and final dispatch
checks detect races. These are not an OS sandbox or an atomic filesystem lock:
the executor must consume only bound_params, immediately after verification.
Bindings and review state are process-local, not cross-process resume tokens.
Each preparation creates a new review attempt; rendering/resuming a stored gate
keeps its ID. Verification is repeatable at engine and adapter boundaries;
exactly-once dispatch remains the runtime's responsibility.
"""

from __future__ import annotations

import copy
import csv
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import stat
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Optional

from .artifact_policy import merge_artifact_policies, validate_artifact_policies
from .execution_mode import ExecutionMode
from .redaction import redact_recursive

_SEAL_KEY = secrets.token_bytes(32)
_MAX_BYTES = 10 * 1024 * 1024
_MAX_TOTAL_BYTES = 32 * 1024 * 1024
_MAX_ROWS = 100000
_MAX_COLUMNS = 256
_MAX_FILES = 32
_MAX_SOURCE_ENTRIES = 256
_MAX_SOURCE_DEPTH = 4
_MAX_ERRORS = 100000
_FORMATS = {".csv", ".json", ".xlsx"}
_EVIDENCE_KEYS = (
    "artifact_id", "version", "content_hash", "manifest_hash", "rule_version",
    "rows", "row_count", "errors", "sample_row_ids", "review_id", "tool_call_id",
    "validation_status", "checked_rules",
)
# The graph may update the display version from its own revision ledger. The
# immutable revision and review_id, not that display counter, authorize bytes.
_SNAPSHOT_KEYS = (
    "artifact_id", "revision", "review_id", "tool_call_id", "required",
    "content_hash", "manifest_hash", "params_hash", "policy_hash", "rule_version",
    "manifest", "files", "stage_dir", "source_root", "source_manifest",
    "artifact_scope", "original_params", "policy", "complete", "valid",
    "rows", "row_count", "errors", "sample", "sample_row_ids", "checked_rows",
    "checked_row_ids", "checked_rules", "validation_status",
    "review_nonce", "run_id", "step_id", "execution_mode",
)


class _ArtifactError(ValueError):
    pass


def _digest(value: Any) -> str:
    raw = json.dumps(value, sort_keys=True, ensure_ascii=True, default=str,
                     allow_nan=False, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _sign(value: dict) -> str:
    return hmac.new(_SEAL_KEY, _digest(value).encode(), hashlib.sha256).hexdigest()


def _snapshot_signature(gate):
    return _sign({key: gate.get(key) for key in _SNAPSHOT_KEYS})


def _binding_context_valid(gate):
    return (all(isinstance(gate.get(key), str) and 0 < len(gate[key]) <= 256 and gate[key].strip()
                for key in ("run_id", "step_id"))
            and isinstance(gate.get("execution_mode"), str)
            and gate.get("execution_mode") in {mode.value for mode in ExecutionMode})


def resolve_artifact_policy(registry, module: str, action: str, company_policy=None,
                            *, artifact_policies=None) -> dict:
    """Resolve trusted rules; company_policy is a legacy keyword alias only.

    registry.artifact_policies optionally holds current host policies, so the
    final adapter boundary can detect host-policy changes without global state.
    Malformed or conflicting policy data makes the gate required and blocked.
    """
    capability = f"{module}.{action}"
    try:
        definition = registry.get_action(module, action)
        if not isinstance(definition, dict):
            raise ValueError("artifact_policy_invalid")
        sources = [definition.get("artifact_policy", {})]
        for mapping in (getattr(registry, "artifact_policies", None), company_policy, artifact_policies):
            sources.append(validate_artifact_policies(mapping).get(capability, {}))
        policy = merge_artifact_policies(*sources)
        if policy["required"]:
            policy.update(capability=capability, risk_level=definition.get("risk_level"),
                          policy_sources_hash=_digest(sources))
            policy.setdefault("rule_version", "artifact-v1")
            policy.setdefault("scope", "input_contract")
        return policy
    except Exception:
        return {"required": True, "capability": capability, "rule_version": "unavailable",
                "scope": "input_contract", "policy_error": "artifact_policy_invalid"}


def _is_link(info):
    return stat.S_ISLNK(info.st_mode) or bool(getattr(info, "st_file_attributes", 0) & 0x400)


def _plain_path(path: Path) -> bool:
    return all(not _is_link(part.lstat()) for part in (path, *path.parents))


def _fingerprint(info):
    return [info.st_dev, info.st_ino, info.st_mode, info.st_size,
            info.st_mtime_ns, info.st_ctime_ns, info.st_nlink]


def _descriptor_identity(info):
    # On Windows, path stat can report creation time as ctime while fstat
    # reports metadata change time. Compare ctime only within the same API.
    fingerprint = _fingerprint(info)
    return fingerprint[:5] + fingerprint[6:] if os.name == "nt" else fingerprint


def _check_name(name):
    if name.startswith("."):
        raise _ArtifactError("hidden_file_forbidden")
    if any(word in name.lower() for word in ("secret", "credential", "token", "password", "api_key")):
        raise _ArtifactError("sensitive_file_forbidden")


def _safe_file(path: Path, root: Path) -> Path:
    if not _plain_path(path):
        raise _ArtifactError("symlink_forbidden")
    resolved = path.resolve(strict=True)
    try:
        relative = resolved.relative_to(root)
    except ValueError:
        raise _ArtifactError("file_outside_root") from None
    for part in relative.parts:
        _check_name(part)
    info = resolved.lstat()
    if not stat.S_ISREG(info.st_mode) or resolved.suffix.lower() not in _FORMATS:
        raise _ArtifactError("unsupported_format")
    if info.st_size > _MAX_BYTES:
        raise _ArtifactError("file_limit")
    return resolved


def _read_checked(path: Path):
    """Read at most the byte bound through one identity-checked descriptor."""
    if not _plain_path(path):
        raise _ArtifactError("symlink_forbidden")
    before = path.lstat()
    if not stat.S_ISREG(before.st_mode):
        raise _ArtifactError("source_entry_type_forbidden")
    if before.st_size > _MAX_BYTES:
        raise _ArtifactError("file_limit")
    flags = os.O_RDONLY | getattr(os, "O_BINARY", 0) | getattr(os, "O_NOFOLLOW", 0)
    fd = os.open(path, flags)
    with os.fdopen(fd, "rb") as stream:
        opened = os.fstat(stream.fileno())
        if _is_link(opened) or _descriptor_identity(opened) != _descriptor_identity(before):
            raise _ArtifactError("artifact_changed_during_read")
        raw = stream.read(_MAX_BYTES + 1)
        if len(raw) > _MAX_BYTES:
            raise _ArtifactError("file_limit")
        after = os.fstat(stream.fileno())
    if (not _plain_path(path) or _fingerprint(opened) != _fingerprint(after)
            or _fingerprint(before) != _fingerprint(path.lstat()) or len(raw) != before.st_size):
        raise _ArtifactError("artifact_changed_during_read")
    return raw, _fingerprint(before)


def _source_directory_manifest(root: Path) -> list:
    """Full bounded metadata membership; never open incidental source files."""
    pending, manifest = [(root, 0)], []
    while pending:
        directory, depth = pending.pop()
        if not _plain_path(directory):
            raise _ArtifactError("symlink_forbidden")
        before = directory.lstat()
        if not stat.S_ISDIR(before.st_mode):
            raise _ArtifactError("source_entry_type_forbidden")
        with os.scandir(directory) as entries:
            for entry in entries:
                if len(manifest) >= _MAX_SOURCE_ENTRIES:
                    raise _ArtifactError("source_entry_limit")
                if depth >= _MAX_SOURCE_DEPTH:
                    raise _ArtifactError("source_depth_limit")
                _check_name(entry.name)
                info = entry.stat(follow_symlinks=False)
                if _is_link(info):
                    raise _ArtifactError("symlink_forbidden")
                is_dir = stat.S_ISDIR(info.st_mode)
                if not is_dir and not stat.S_ISREG(info.st_mode):
                    raise _ArtifactError("source_entry_type_forbidden")
                path = Path(entry.path)
                manifest.append({"entry_id": "entry-" + _digest(path.relative_to(root).as_posix()),
                                 "kind": "directory" if is_dir else "file",
                                 "size": None if is_dir else info.st_size})
                if is_dir:
                    pending.append((path, depth + 1))
        if not _plain_path(directory) or _fingerprint(before) != _fingerprint(directory.lstat()):
            raise _ArtifactError("source_membership_changed")
    return sorted(manifest, key=lambda item: item["entry_id"])


def _headers(headers):
    if (not headers or len(headers) > _MAX_COLUMNS
            or any(not isinstance(h, str) or not h.strip() for h in headers)
            or len(set(headers)) != len(headers)):
        raise _ArtifactError("invalid_columns")


def _json_pairs(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise _ArtifactError("duplicate_columns")
        result[key] = value
    return result


def _invalid_constant(value):
    raise _ArtifactError("invalid_rows")


def _read_rows(raw: bytes, suffix: str, file_id: str, policy: dict) -> list:
    values = []

    def append(row):
        if len(values) >= _MAX_ROWS:
            raise _ArtifactError("row_limit")
        if not isinstance(row, dict) or len(row) > _MAX_COLUMNS or any(not isinstance(k, str) for k in row):
            raise _ArtifactError("invalid_rows")
        values.append({**row, "row_id": f"{file_id}:row-{len(values) + 1:05d}"})

    if suffix == ".csv":
        reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")), strict=True)
        _headers(reader.fieldnames or [])
        for row in reader:
            if None in row or any(value is None for value in row.values()):
                raise _ArtifactError("invalid_rows")
            append(row)
    elif suffix == ".json":
        parsed = json.loads(raw, object_pairs_hook=_json_pairs, parse_constant=_invalid_constant)
        if not isinstance(parsed, list):
            raise _ArtifactError("rows_array_required")
        for row in parsed:
            append(row)
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise _ArtifactError("xlsx_reader_unavailable") from None
        # Bound the expanded container before invoking the optional XML reader.
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 1024 or sum(item.file_size for item in entries) > _MAX_TOTAL_BYTES:
                raise _ArtifactError("file_limit")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            sheets = [workbook[policy["sheet"]]] if policy.get("sheet") else workbook.worksheets
            for sheet in sheets:
                if (sheet.max_column or 0) > _MAX_COLUMNS or (sheet.max_row or 0) > _MAX_ROWS + 1:
                    raise _ArtifactError("row_limit")
                iterator = sheet.iter_rows(values_only=True)
                headers = list(next(iterator, ()))
                _headers(headers)
                for cells in iterator:
                    append(dict(zip(headers, cells)))
        finally:
            workbook.close()
    else:
        raise _ArtifactError("unsupported_format")
    return values


def select_review_sample(rows: list, errors: list, seed: str) -> list:
    bad = {error.get("row_id") for error in errors}
    good = sorted((row for row in rows if row["row_id"] not in bad),
                  key=lambda row: _digest([seed, row["row_id"]]))
    return good[:5] + [row for row in rows if row["row_id"] in bad]


def _validation_result(result, rows, policy):
    if (not isinstance(result, dict) or type(result.get("complete")) is not bool
            or type(result.get("valid")) is not bool or not isinstance(result.get("errors"), list)
            or len(result["errors"]) > _MAX_ERRORS):
        raise _ArtifactError("artifact_validator_invalid_result")
    count = result.get("row_count", result.get("rows"))
    if type(count) is not int or count != len(rows):
        raise _ArtifactError("artifact_validator_invalid_result")
    if "rows" in result and (type(result["rows"]) is not int or result["rows"] != count):
        raise _ArtifactError("artifact_validator_invalid_result")
    ids = {row["row_id"] for row in rows}
    errors = []
    public_fields = {"rows", "file", "rules", *(policy.get("columns") or [])}
    for error in result["errors"]:
        if not isinstance(error, dict):
            raise _ArtifactError("artifact_validator_invalid_result")
        row_id, code = error.get("row_id"), error.get("code")
        if ((row_id is not None and (not isinstance(row_id, str) or row_id not in ids))
                or not isinstance(code, str) or not re.fullmatch(r"[a-z][a-z0-9_]{0,79}", code)):
            raise _ArtifactError("artifact_validator_invalid_result")
        field = error.get("field")
        # Plugin exception messages and arbitrary values never become evidence.
        if not isinstance(field, str) or field not in public_fields:
            field = "content"
        errors.append({"row_id": row_id, "field": field, "code": code})
    checked = result.get("checked_rules", [])
    if (not isinstance(checked, list) or len(checked) > 256
            or any(not isinstance(x, str) or not re.fullmatch(r"[A-Za-z0-9_][A-Za-z0-9_.-]{0,79}", x)
                   for x in checked)):
        raise _ArtifactError("artifact_validator_invalid_result")
    if result["complete"] and not set(policy.get("required_rules", [])).issubset(checked):
        raise _ArtifactError("artifact_required_rules_unchecked")
    if result["valid"] and (not result["complete"] or errors):
        raise _ArtifactError("artifact_validator_invalid_result")
    if not result["valid"] and not errors:
        errors = [{"row_id": None, "field": "content", "code": "artifact_validation_failed"}]
    declared_rules = set(policy.get("required_rules", [])) | set(policy.get("rules", {}))
    return {"complete": result["complete"], "valid": result["valid"], "errors": errors,
            "rows": count, "row_count": count, "checked_rules": sorted(set(checked) & declared_rules),
            "validation_status": "complete" if result["complete"] else "unavailable"}


def _revision_files_match(gate):
    root, stage = Path(gate["source_root"]), Path(gate["stage_dir"])
    if not _plain_path(root) or not _plain_path(stage) or stage.resolve().is_relative_to(root):
        return False
    directory_scope = gate["artifact_scope"] == "directory_root"
    if directory_scope and _source_directory_manifest(root) != gate["source_manifest"]:
        return False
    files = gate["files"]
    if not files or len(files) > _MAX_FILES:
        return False
    expected = {Path(item["staged"]) for item in files}
    if (len(expected) != len(files) or any(path.parent != stage for path in expected)
            or set(stage.iterdir()) != expected):
        return False
    for item in files:
        source = _safe_file(Path(item["source"]), root)
        for key, path in (("source", source), ("staged", Path(item["staged"]))):
            raw, fingerprint = _read_checked(path)
            if hashlib.sha256(raw).hexdigest() != item["hash"]:
                return False
            if fingerprint != item[key + "_state"]:
                return False
    return (set(stage.iterdir()) == expected and _plain_path(stage) and _plain_path(root)
            and (not directory_scope or _source_directory_manifest(root) == gate["source_manifest"]))


def prepare_artifact_review(params, artifacts, policy, *, run_id, tool_call_id,
                            staging_root=None, previous=None, registry=None, validator=None,
                            step_id=None, execution_mode=None) -> dict:
    """Freeze real files and invoke only an explicitly supplied content validator.

    snapshot has rows (list), row_count, manifest, artifact_id, content_hash,
    manifest_hash and artifact_scope. A private copy is passed to the plugin;
    plugin changes cannot alter the engine's byte identity or review contents.
    step_id and execution_mode come from trusted execution state, not params.
    Omission permits content inspection but cannot authorize execution.
    """
    identity = "artifact-" + _digest([run_id, params.get("root_dir") if isinstance(params, dict) else None])[:24]
    gate = {"required": True, "artifact_id": identity, "version": 1, "content_hash": None,
            "manifest_hash": None, "rule_version": None, "rows": 0, "row_count": 0,
            "errors": [], "sample_row_ids": [], "review_id": None, "tool_call_id": tool_call_id,
            "complete": False, "valid": False, "decision": None, "files": [], "sample": [],
            "checked_rows": [], "checked_row_ids": [], "checked_rules": [],
            "validation_status": "unavailable", "run_id": run_id, "step_id": step_id,
            "execution_mode": None, "review_nonce": secrets.token_hex(16)}
    created = []
    try:
        if not isinstance(params, dict) or not isinstance(policy, dict):
            raise _ArtifactError("artifact_policy_invalid")
        if policy.get("policy_error"):
            raise _ArtifactError("artifact_policy_invalid")
        if execution_mode is not None:
            try:
                gate["execution_mode"] = ExecutionMode.coerce(execution_mode).value
            except (TypeError, ValueError):
                raise _ArtifactError("artifact_binding_context_invalid") from None
        # Check resolved-policy extensions separately from the generic schema.
        generic = {key: value for key, value in policy.items()
                   if key not in {"capability", "risk_level", "policy_sources_hash"}}
        try:
            merge_artifact_policies(generic)
        except (TypeError, ValueError):
            raise _ArtifactError("artifact_policy_invalid") from None
        gate.update(artifact_id="artifact-" + _digest([run_id, params.get("root_dir")])[:24],
                    params_hash=_digest(params), policy_hash=_digest(policy),
                    rule_version=str(policy.get("rule_version", "artifact-v1")),
                    policy=copy.deepcopy(policy), original_params=copy.deepcopy(params),
                    capability=policy.get("capability"))
        validator = validator if validator is not None else getattr(registry, "artifact_validator", None)
        if not callable(getattr(validator, "validate", None)):
            raise _ArtifactError("artifact_validator_unavailable")
        root_value = params.get("root_dir")
        if not isinstance(root_value, str) or not root_value or "://" in root_value:
            raise _ArtifactError("explicit_root_required")
        raw_root = Path(root_value)
        if not _plain_path(raw_root):
            raise _ArtifactError("symlink_forbidden")
        root = raw_root.resolve(strict=True)
        if not root.is_dir():
            raise _ArtifactError("explicit_root_required")
        refs = params.get("artifact_files")
        if refs is None:
            refs = [item.get("path") for item in artifacts or [] if isinstance(item, dict)
                    and item.get("path") and item.get("type") not in
                    {"mock_artifact", "output_root_dir", "activity_dir"}]
        if not isinstance(refs, list) or not refs or len(refs) > _MAX_FILES:
            raise _ArtifactError("explicit_files_required")
        paths = []
        for ref in refs:
            if not isinstance(ref, str) or not ref or "://" in ref:
                raise _ArtifactError("actual_file_required")
            path = Path(ref)
            paths.append(_safe_file(path if path.is_absolute() else root / path, root))
        if len(set(paths)) != len(paths):
            raise _ArtifactError("duplicate_file")
        scope = policy.get("scope", "input_contract")
        if scope == "input_contract":
            scope = "explicit_file_set" if params.get("artifact_files") is not None else "directory_root"
        if scope not in {"explicit_file_set", "directory_root"}:
            raise _ArtifactError("invalid_artifact_scope")
        gate.update(artifact_scope=scope, source_root=str(root), source_manifest=[])
        if scope == "directory_root":
            gate["source_manifest"] = _source_directory_manifest(root)
        # Frozen files are outside the source tree even for explicit bundles.
        stage_parent = Path(staging_root or tempfile.gettempdir())
        for parent in (stage_parent, *stage_parent.parents):
            if parent.exists() and not _plain_path(parent):
                raise _ArtifactError("symlink_forbidden")
        if stage_parent.resolve().is_relative_to(root):
            stage_parent = Path(tempfile.gettempdir())
        if stage_parent.resolve().is_relative_to(root):
            raise _ArtifactError("staging_outside_source_required")
        for parent in (stage_parent, *stage_parent.parents):
            if parent.exists() and not _plain_path(parent):
                raise _ArtifactError("symlink_forbidden")
        stage_parent.mkdir(parents=True, exist_ok=True)
        if not _plain_path(stage_parent):
            raise _ArtifactError("symlink_forbidden")
        rows, blobs, manifest, total_bytes = [], [], [], 0
        for index, path in enumerate(sorted(paths, key=lambda p: p.relative_to(root).as_posix())):
            raw, fingerprint = _read_checked(path)
            total_bytes += len(raw)
            if total_bytes > _MAX_TOTAL_BYTES:
                raise _ArtifactError("total_file_limit")
            file_id = f"file-{index + 1:03d}"
            digest = hashlib.sha256(raw).hexdigest()
            manifest.append({"file_id": file_id, "entry_id": "entry-" + _digest(path.relative_to(root).as_posix()),
                             "hash": digest, "size": len(raw), "format": path.suffix.lower()})
            blobs.append((path, raw, fingerprint, file_id))
            rows.extend(_read_rows(raw, path.suffix.lower(), file_id, policy))
            if len(rows) > _MAX_ROWS:
                raise _ArtifactError("row_limit")
        gate.update(content_hash=_digest([item["hash"] for item in manifest]), manifest=manifest)
        gate["manifest_hash"] = _digest({"files": manifest, "scope": scope,
                                         "source_membership": gate["source_manifest"]})
        gate["revision"] = _digest([gate["manifest_hash"], gate["params_hash"], gate["policy_hash"]])
        if previous and previous.get("artifact_id", gate["artifact_id"]) == gate["artifact_id"]:
            version = previous.get("version", 1)
            if type(version) is not int or version < 1:
                raise _ArtifactError("artifact_previous_invalid")
            gate["version"] = version + int(previous.get("revision") != gate["revision"])
        # Content identity is reusable; permission for a new execution is not.
        # A per-attempt nonce also separates same-call retries and A -> B -> A.
        gate["review_id"] = "review-" + _digest([
            run_id, step_id, tool_call_id, gate["execution_mode"], gate["revision"], gate["review_nonce"],
        ])[:24]
        snapshot = {key: copy.deepcopy(gate[key]) for key in
                    ("artifact_id", "manifest", "content_hash", "manifest_hash", "artifact_scope")}
        snapshot.update(rows=copy.deepcopy(rows), row_count=len(rows))
        try:
            result = validator.validate(snapshot, copy.deepcopy(policy))
        except Exception:
            raise _ArtifactError("artifact_validator_failed") from None
        gate.update(_validation_result(result, rows, policy))
        columns = policy.get("columns") or []
        projected = [{"row_id": row["row_id"], **{key: row.get(key) for key in columns if key != "row_id"}}
                     for row in rows]
        # Optional readers can return dates and other scalars. Public evidence
        # stays JSON-compatible without changing the validator's parsed inputs.
        projected = json.loads(json.dumps(projected, default=str, allow_nan=False))
        gate["checked_rows"], _ = redact_recursive(projected, max_list_items=_MAX_ROWS)
        gate["checked_row_ids"] = [row["row_id"] for row in rows]
        gate["sample"] = select_review_sample(gate["checked_rows"], gate["errors"], gate["content_hash"])
        gate["sample_row_ids"] = [row["row_id"] for row in gate["sample"]]
        stage = Path(tempfile.mkdtemp(prefix="artifact-", dir=stage_parent)).resolve()
        gate["stage_dir"] = str(stage)
        for path, raw, fingerprint, file_id in blobs:
            staged = stage / (file_id + path.suffix.lower())
            created.append(staged)
            with staged.open("xb") as stream:
                stream.write(raw)
            frozen, staged_state = _read_checked(staged)
            if frozen != raw:
                raise _ArtifactError("artifact_changed_during_read")
            gate["files"].append({"source": str(path), "staged": str(staged),
                                  "hash": hashlib.sha256(raw).hexdigest(),
                                  "source_state": fingerprint, "staged_state": staged_state})
        if scope == "directory_root" and _source_directory_manifest(root) != gate["source_manifest"]:
            raise _ArtifactError("source_membership_changed")
        if not _revision_files_match(gate):
            raise _ArtifactError("artifact_changed_during_read")
        gate["snapshot_seal"] = _snapshot_signature(gate)
        return gate
    except Exception as exc:
        code = str(exc) if isinstance(exc, _ArtifactError) else "artifact_unavailable"
        gate.update(complete=False, valid=False, validation_status="unavailable",
                    errors=[{"row_id": None, "field": "file", "code": code}])
        # Only remove files this invocation created, never recursively delete a
        # supplied directory or an externally changed staging parent.
        stage = Path(gate["stage_dir"]) if gate.get("stage_dir") else None
        if stage is not None:
            try:
                if _plain_path(stage):
                    for path in created:
                        if path.parent == stage:
                            path.unlink(missing_ok=True)
                    stage.rmdir()
            except OSError:
                pass
        gate.pop("snapshot_seal", None)
        return gate


def artifact_event(kind: str, gate: dict, **extra) -> dict:
    event = {"type": kind, **{key: gate.get(key) for key in _EVIDENCE_KEYS},
             "manifest": gate.get("manifest"), "checked_rows": gate.get("checked_rows"),
             "checked_row_ids": gate.get("checked_row_ids"), "artifact_scope": gate.get("artifact_scope"),
             "source_manifest": gate.get("source_manifest"),
             "source": "frozen-" + gate.get("artifact_scope", "file-manifest").replace("_", "-"),
             "complete": gate.get("complete", False), "valid": gate.get("valid", False)}
    if kind == "artifact.review":
        event.update({key: gate[key] for key in ("decision", "comment") if key in gate})
    event.update({key: value for key, value in extra.items()
                  if key in {"decision", "comment", "error_type", "sample"}})
    event = redact_recursive(event, max_list_items=max(_MAX_ROWS, _MAX_ERRORS))[0]
    if "comment" in event:
        event["comment"] = event["comment"][:4000] if isinstance(event["comment"], str) else ""
    return event


def review_interaction(gate: dict) -> dict:
    return {"interaction_id": gate.get("review_id"), "type": "artifact_review", "status": "pending",
            "question": "Review the checked artifact revision", "reason": "Execution requires artifact review",
            "actions": [], "can_approve": bool(gate.get("complete") and gate.get("valid")
                                                and not gate.get("errors") and gate.get("snapshot_seal")
                                                and _binding_context_valid(gate)),
            "evidence": artifact_event("artifact.review", gate), "sample": copy.deepcopy(gate.get("sample", [])),
            "decisions": ["approve", "reject", "request_changes"]}


def artifact_revision_matches(gate: dict, params: dict, policy: dict) -> bool:
    try:
        return bool(not policy.get("policy_error")
                    and hmac.compare_digest(str(gate.get("snapshot_seal", "")), _snapshot_signature(gate))
                    and gate.get("params_hash") == _digest(params)
                    and gate.get("policy_hash") == _digest(policy)
                    and _revision_files_match(gate))
    except Exception:
        return False


def bind_artifact_review(gate: dict, params: dict, policy: dict) -> dict:
    bound = copy.deepcopy(gate)
    for key in ("seal", "bound_params"):
        bound.pop(key, None)
    if (gate.get("decision") != "approve" or not gate.get("complete") or not gate.get("valid")
            or gate.get("errors") or gate.get("error_type")):
        bound["error_type"] = "artifact_review_required"
    elif not _binding_context_valid(gate):
        bound["error_type"] = "artifact_binding_context_required"
    elif not artifact_revision_matches(gate, params, policy):
        bound.update(error_type="artifact_approval_invalidated", decision=None)
    elif policy.get("capability") and gate.get("capability") != policy["capability"]:
        bound.update(error_type="artifact_approval_invalidated", decision=None)
    else:
        bound["bound_params"] = {**copy.deepcopy(params), "root_dir": gate["stage_dir"],
                                 "artifact_files": [item["staged"] for item in gate["files"]]}
        bound["seal"] = _sign(bound)
    return bound


def verify_adapter_binding(adapter_input, registry, *, run_id=None) -> Optional[str]:
    """Verify against current run/step/mode, independently of binding metadata.

    run_id may be supplied explicitly or via adapter_input.run_id. Missing
    current context blocks. Both forms, when present, must agree. This function
    does not consume approvals because dispatch performs two boundary checks.
    """
    binding = getattr(adapter_input, "artifact_binding", None)
    policy = resolve_artifact_policy(registry, adapter_input.module, adapter_input.action)
    if not binding:
        return "artifact_review_runtime_required" if policy.get("required") else None
    if not isinstance(binding, dict) or not binding.get("seal"):
        return "artifact_review_runtime_required"
    try:
        content = {key: value for key, value in binding.items() if key != "seal"}
        if not hmac.compare_digest(str(binding["seal"]), _sign(content)):
            return "artifact_binding_invalid"
        if binding.get("capability") != f"{adapter_input.module}.{adapter_input.action}":
            return "artifact_binding_invalid"
        adapter_run_id = getattr(adapter_input, "run_id", None)
        current_run_id = run_id if run_id is not None else adapter_run_id
        mode = getattr(adapter_input, "execution_mode", None)
        context = {"run_id": current_run_id, "step_id": getattr(adapter_input, "step_id", None),
                   "execution_mode": ExecutionMode.coerce(mode).value if mode is not None else None}
        if not _binding_context_valid(binding) or not _binding_context_valid(context):
            return "artifact_binding_context_required"
        if adapter_run_id is not None and adapter_run_id != current_run_id:
            return "artifact_approval_invalidated"
        if any(binding.get(key) != value for key, value in context.items()):
            return "artifact_approval_invalidated"
        # Legacy graph metadata remains supported, but only after authentication.
        policy = resolve_artifact_policy(registry, adapter_input.module, adapter_input.action,
                                         binding.get("company_policy"),
                                         artifact_policies=binding.get("artifact_policies"))
        if (binding.get("decision") != "approve" or binding.get("error_type")
                or not binding.get("complete") or not binding.get("valid") or binding.get("errors")):
            return "artifact_binding_invalid"
        if adapter_input.params != binding.get("bound_params"):
            return "artifact_approval_invalidated"
        if not artifact_revision_matches(binding, binding["original_params"], policy):
            return "artifact_approval_invalidated"
        return None
    except Exception:
        return "artifact_binding_invalid"


def artifact_blocked_observation(step_id: str, error_type: str, gate=None):
    from .schemas import Observation
    if not gate:
        gate = {"artifact_id": "artifact-" + _digest([step_id])[:24], "tool_call_id": step_id,
                "rows": 0, "row_count": 0, "complete": False, "valid": False,
                "errors": [{"row_id": None, "field": "runtime", "code": error_type}]}
    artifacts = [artifact_event("artifact.check", gate)]
    summary = error_type
    if gate.get("decision") in {"approve", "reject", "request_changes", "invalidated"}:
        feedback = artifact_event("artifact.review", gate)
        artifacts.append(feedback)
        # Preserve revision feedback in both the next full tool Observation and
        # its compact summary; stable error codes remain separate from comments.
        summary += "; review_decision=" + feedback["decision"]
        if feedback.get("comment"):
            summary += ": " + feedback["comment"]
    return Observation(step_id=step_id, status="failed", summary=summary,
                       error={"type": error_type, "message": error_type, "retryable": False},
                       artifacts=artifacts,
                       suggested_next_action="revise_artifact_and_request_review")
